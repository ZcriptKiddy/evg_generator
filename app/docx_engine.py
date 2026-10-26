from __future__ import annotations

import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageOps, UnidentifiedImageError


MAX_DOCX_BYTES = 7 * 1024 * 1024
IMAGE_HEIGHT_CM = 5.3
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}
PLACEHOLDER_RE = re.compile(r"{{\s*([\w]+)\s*}}", re.UNICODE)


class TemplateMissingError(RuntimeError):
    pass


class ImageMissingError(RuntimeError):
    def __init__(self, missing: list[str]) -> None:
        self.missing = missing
        super().__init__("No se encontraron imagenes: " + ", ".join(missing))


class DocumentTooLargeError(RuntimeError):
    pass


@dataclass(frozen=True)
class RenderStats:
    size_bytes: int
    image_quality: int
    image_max_side: int
    image_count: int


@dataclass(frozen=True)
class ImageAttempt:
    max_side: int
    quality: int


IMAGE_ATTEMPTS = (
    ImageAttempt(max_side=1600, quality=86),
    ImageAttempt(max_side=1300, quality=82),
    ImageAttempt(max_side=1050, quality=78),
    ImageAttempt(max_side=850, quality=74),
    ImageAttempt(max_side=680, quality=68),
    ImageAttempt(max_side=520, quality=62),
)


def render_docx(
    *,
    template_path: Path,
    payload: dict[str, Any],
    output_path: Path,
    work_dir: Path,
    images_zip_path: Path | None = None,
    images_dir_path: Path | None = None,
) -> RenderStats:
    if not template_path.exists():
        raise TemplateMissingError(f"Falta la plantilla: {template_path.name}")

    extracted_dir = work_dir / "imagenes_extraidas"
    image_index = _extract_zip_images(images_zip_path, extracted_dir) if images_zip_path else {}
    if images_dir_path:
        image_index.update(_index_image_files(images_dir_path))
    text_values = _text_context(payload)
    image_values = _image_context(payload, image_index)

    missing = [ref for refs in image_values.missing.values() for ref in refs]
    if missing:
        raise ImageMissingError(sorted(set(missing)))

    last_size = 0
    for attempt in IMAGE_ATTEMPTS:
        compressed_dir = work_dir / f"imagenes_{attempt.max_side}_{attempt.quality}"
        compressed_images = _compress_image_context(
            image_values.resolved,
            compressed_dir,
            attempt=attempt,
        )

        doc = Document(str(template_path))
        _replace_everywhere(doc, text_values, compressed_images)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))

        last_size = output_path.stat().st_size
        image_count = sum(len(paths) for paths in compressed_images.values())
        if last_size <= MAX_DOCX_BYTES:
            return RenderStats(
                size_bytes=last_size,
                image_quality=attempt.quality,
                image_max_side=attempt.max_side,
                image_count=image_count,
            )

    raise DocumentTooLargeError(
        f"El documento quedo en {last_size / 1024 / 1024:.2f} MB aun despues de comprimir imagenes."
    )


def analyze_docx_placeholders(template_path: Path) -> dict[str, list[str]]:
    if not template_path.exists():
        raise TemplateMissingError(f"Falta la plantilla: {template_path.name}")

    doc = Document(str(template_path))
    placeholders = {"txt": [], "img": []}
    seen = {"txt": set(), "img": set()}

    for paragraph in _iter_document_paragraphs(doc):
        for match in PLACEHOLDER_RE.finditer(paragraph.text):
            key = match.group(1)
            kind = ""
            if key.startswith("Txt_"):
                kind = "txt"
            elif key.startswith("Img_"):
                kind = "img"

            if kind and key not in seen[kind]:
                seen[kind].add(key)
                placeholders[kind].append(key)

    return placeholders


@dataclass
class ImageContext:
    resolved: dict[str, list[Path]]
    missing: dict[str, list[str]]


def _text_context(payload: dict[str, Any]) -> dict[str, str]:
    values: dict[str, str] = {}
    for key, value in payload.items():
        if key.startswith("Img_"):
            continue
        if isinstance(value, (str, int, float, bool)) or value is None:
            values[key] = "" if value is None else str(value)
    return values


def _image_context(payload: dict[str, Any], image_index: dict[str, Path]) -> ImageContext:
    resolved: dict[str, list[Path]] = {}
    missing: dict[str, list[str]] = {}

    for key, value in payload.items():
        if not key.startswith("Img_"):
            continue

        refs = _normalize_image_refs(value)
        paths: list[Path] = []
        absent: list[str] = []

        for ref in refs:
            found = _resolve_image(ref, image_index)
            if found is None:
                absent.append(ref)
                continue

            paths.extend(_with_numbered_siblings(found, image_index))

        if paths:
            resolved[key] = _unique_paths(paths)
        if absent:
            missing[key] = absent

    _merge_numbered_json_keys(resolved)
    return ImageContext(resolved=resolved, missing=missing)


def _normalize_image_refs(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [item for item in value if isinstance(item, str)]
    return []


def _extract_zip_images(zip_path: Path, target_dir: Path) -> dict[str, Path]:
    target_dir.mkdir(parents=True, exist_ok=True)
    image_index: dict[str, Path] = {}

    with zipfile.ZipFile(zip_path) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue

            source_name = Path(info.filename.replace("\\", "/")).name
            ext = Path(source_name).suffix.lower()
            if ext not in IMAGE_EXTENSIONS:
                continue

            safe_name = _unique_file_name(target_dir, source_name)
            target_path = target_dir / safe_name
            with archive.open(info) as src, target_path.open("wb") as dst:
                shutil.copyfileobj(src, dst)

            image_index[target_path.name.lower()] = target_path
            image_index[target_path.stem.lower()] = target_path

    return image_index


def _index_image_files(images_dir: Path) -> dict[str, Path]:
    image_index: dict[str, Path] = {}
    if not images_dir.exists():
        return image_index

    for path in images_dir.iterdir():
        if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        image_index[path.name.lower()] = path
        image_index[path.stem.lower()] = path

    return image_index


def _unique_file_name(directory: Path, file_name: str) -> str:
    candidate = Path(file_name).name
    stem = Path(candidate).stem
    suffix = Path(candidate).suffix
    counter = 1

    while (directory / candidate).exists():
        candidate = f"{stem}_{counter}{suffix}"
        counter += 1

    return candidate


def _resolve_image(ref: str, image_index: dict[str, Path]) -> Path | None:
    name = Path(ref.replace("\\", "/")).name.lower()
    stem = Path(name).stem.lower()
    return image_index.get(name) or image_index.get(stem)


def _with_numbered_siblings(primary: Path, image_index: dict[str, Path]) -> list[Path]:
    stem = primary.stem.lower()
    pattern = re.compile(rf"^{re.escape(stem)}_[0-9]+$", re.IGNORECASE)
    siblings = [
        path
        for key, path in image_index.items()
        if Path(key).suffix and pattern.match(path.stem.lower())
    ]
    return [primary, *_natural_path_sort(siblings)]


def _natural_path_sort(paths: list[Path]) -> list[Path]:
    def key(path: Path) -> list[Any]:
        return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", path.stem)]

    return sorted(paths, key=key)


def _unique_paths(paths: list[Path]) -> list[Path]:
    unique: list[Path] = []
    seen: set[Path] = set()
    for path in paths:
        resolved = path.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique.append(path)
    return unique


def _merge_numbered_json_keys(resolved: dict[str, list[Path]]) -> None:
    for key in list(resolved):
        match = re.match(r"^(Img_.+?)_[0-9]+$", key)
        if not match:
            continue

        base_key = match.group(1)
        if base_key not in resolved:
            continue

        resolved[base_key] = _unique_paths([*resolved[base_key], *resolved[key]])


def _compress_image_context(
    image_values: dict[str, list[Path]],
    target_dir: Path,
    *,
    attempt: ImageAttempt,
) -> dict[str, list[Path]]:
    target_dir.mkdir(parents=True, exist_ok=True)
    cache: dict[Path, Path] = {}
    compressed: dict[str, list[Path]] = {}

    for key, paths in image_values.items():
        compressed[key] = []
        for path in paths:
            resolved = path.resolve()
            if resolved not in cache:
                cache[resolved] = _compress_image(path, target_dir, attempt=attempt)
            compressed[key].append(cache[resolved])

    return compressed


def _compress_image(source: Path, target_dir: Path, *, attempt: ImageAttempt) -> Path:
    try:
        with Image.open(source) as image:
            image = ImageOps.exif_transpose(image)
            image.thumbnail((attempt.max_side, attempt.max_side), Image.Resampling.LANCZOS)

            if image.mode in ("RGBA", "LA"):
                background = Image.new("RGB", image.size, "white")
                background.paste(image, mask=image.getchannel("A"))
                image = background
            elif image.mode != "RGB":
                image = image.convert("RGB")

            target_path = target_dir / f"{source.stem}_{uuid4().hex[:8]}.jpg"
            image.save(
                target_path,
                "JPEG",
                quality=attempt.quality,
                optimize=True,
                progressive=True,
            )
            return target_path
    except UnidentifiedImageError as exc:
        raise RuntimeError(f"No pude leer la imagen {source.name}") from exc


def _replace_everywhere(
    doc: Document,
    text_values: dict[str, str],
    image_values: dict[str, list[Path]],
) -> None:
    for paragraph in _iter_document_paragraphs(doc):
        _replace_paragraph(paragraph, text_values, image_values)


def _iter_document_paragraphs(doc: Document):
    for section in doc.sections:
        header_containers = (
            section.header,
            section.first_page_header,
            section.even_page_header,
        )
        for container in header_containers:
            for paragraph in _iter_paragraphs(container):
                yield paragraph

    for paragraph in _iter_paragraphs(doc):
        yield paragraph

    for section in doc.sections:
        footer_containers = (
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        )
        for container in footer_containers:
            for paragraph in _iter_paragraphs(container):
                yield paragraph


def _iter_paragraphs(container: Any):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph

    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _replace_paragraph(paragraph: Any, text_values: dict[str, str], image_values: dict[str, list[Path]]) -> None:
    original = paragraph.text
    if "{{" not in original:
        return

    matches = list(PLACEHOLDER_RE.finditer(original))
    if not matches:
        return

    has_images = any(match.group(1) in image_values for match in matches)
    if not has_images:
        replaced = PLACEHOLDER_RE.sub(lambda match: text_values.get(match.group(1), match.group(0)), original)
        if replaced != original:
            _set_paragraph_text(paragraph, replaced)
        return

    _clear_paragraph(paragraph)
    cursor = 0

    for match in matches:
        before = original[cursor : match.start()]
        if before:
            paragraph.add_run(_replace_text_tokens(before, text_values))

        key = match.group(1)
        if key in image_values:
            _add_images(paragraph, image_values[key])
        else:
            paragraph.add_run(text_values.get(key, match.group(0)))

        cursor = match.end()

    after = original[cursor:]
    if after:
        paragraph.add_run(_replace_text_tokens(after, text_values))


def _replace_text_tokens(text: str, text_values: dict[str, str]) -> str:
    return PLACEHOLDER_RE.sub(lambda match: text_values.get(match.group(1), match.group(0)), text)


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _clear_paragraph(paragraph: Any) -> None:
    for run in paragraph.runs:
        run.text = ""


def _add_images(paragraph: Any, paths: list[Path]) -> None:
    for index, path in enumerate(paths):
        if index:
            paragraph.add_run(" ")
        run = paragraph.add_run()
        run.add_picture(str(path), height=Cm(IMAGE_HEIGHT_CM))
